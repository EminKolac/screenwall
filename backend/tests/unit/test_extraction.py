from app.extraction.base import BlockType
from app.extraction.docx import DocxExtractor
from app.extraction.pdf import PdfExtractor
from app.extraction.xlsx import XlsxExtractor
from app.models.document import FileKind


def test_docx_structure(docx_en):
    c = DocxExtractor().extract(docx_en, "a.docx")
    assert c.kind == FileKind.docx
    assert any(b.type == BlockType.heading for b in c.blocks)
    table = next(b for b in c.blocks if b.type == BlockType.table)
    assert any(cell.text == "Name" for cell in table.cells)
    assert all(b.block_id for b in c.blocks)  # stable ids assigned
    assert "England" in c.plain_text and "John Smith" in c.plain_text


def test_xlsx_cells_with_addresses(xlsx_tr):
    c = XlsxExtractor().extract(xlsx_tr, "a.xlsx")
    tables = [b for b in c.blocks if b.type == BlockType.table]  # sheet-name heading precedes it
    assert tables
    cells = tables[0].cells
    assert any(cell.address == "A1" for cell in cells)
    assert tables[0].sheet
    assert "Ahmet Yılmaz" in c.plain_text


def test_xlsx_block_ids_and_locations_never_carry_the_sheet_name():
    """Block ids become unit ids in layer 3 (anonymized.json, edits.json) and layer 4
    (findings.json), and the location reaches the review screen: none of them is masked. The raw
    sheet title may live only where the engine masks it (the heading text) or clears it
    (`Block.sheet`)."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Ahmet Yılmaz"
    wb.active.append(["Ad"])
    wb.active.append(["Ahmet Yılmaz"])
    wb.create_sheet("Boş sayfa")  # no cells: emits nothing, but keeps its position
    wb.create_sheet("Ayşe Kaya").append(["Maaş"])
    buf = io.BytesIO()
    wb.save(buf)

    c = XlsxExtractor().extract(buf.getvalue(), "x.xlsx")
    sheets = [b for b in c.blocks if b.location != "comment"]
    assert [(b.block_id, b.type, b.location) for b in sheets] == [
        ("x1-h", BlockType.heading, "Sayfa 1"), ("x1-t", BlockType.table, "Sayfa 1"),
        ("x3-h", BlockType.heading, "Sayfa 3"), ("x3-t", BlockType.table, "Sayfa 3"),
    ]
    assert [b.text for b in sheets if b.type == BlockType.heading] == ["Ahmet Yılmaz", "Ayşe Kaya"]
    assert [b.sheet for b in sheets if b.type == BlockType.table] == ["Ahmet Yılmaz", "Ayşe Kaya"]


def test_pdf_text(pdf_en):
    c = PdfExtractor().extract(pdf_en, "a.pdf")
    assert c.kind == FileKind.pdf
    assert any(b.page == 1 for b in c.blocks)
    assert "finance team" in c.plain_text


# ---------------------------------------------------------------------------------------------
# Kapsam boşlukları — ikisi de probla bulundu, ikisi de sessizdi (uyarı üretmiyordu).
# ---------------------------------------------------------------------------------------------

def test_sticky_note_annotation_content_is_extracted():
    """Text (yapışkan not) annotation'ının içeriği sayfada ÇİZİLMEZ; yalnız /Contents alanında
    durur. `get_text()` onu göremiyordu, dolayısıyla bir yoruma yazılmış IBAN hiç maskelenmiyor
    ve denetçiye de hiç ulaşmıyordu."""
    import fitz

    from app.extraction.dispatcher import extract

    d = fitz.open()
    p = d.new_page()
    p.insert_text((36, 60), "SOZLESME TASLAGI", fontsize=12)
    a = p.add_text_annot((300, 100), "IBAN: TR330006100519786457841326")
    a.set_info(title="Av. Burak Sahin", content="IBAN: TR330006100519786457841326")
    a.update()
    data = d.tobytes()
    d.close()

    _, content = extract(data, "x.pdf", max_bytes=5 * 1024 * 1024)
    assert "TR330006100519786457841326" in content.plain_text
    assert "Burak Sahin" in content.plain_text  # yazar adı da kişisel veri


def test_image_heavy_page_with_header_text_warns_and_blocks_auto_approval():
    """Eski davranışta fail-OPEN: sayfanın büyük bölümü görüntü olsa bile metin katmanı 50
    karakteri geçtiğinde OCR atlanıyor ve HİÇBİR uyarı üretilmiyordu — görüntüdeki PII sessizce
    okunmadan geçip belge temiz rotadan ilerliyordu."""
    import fitz

    from app.config import Settings
    from app.extraction.base import IMAGE_TEXT_UNREAD
    from app.extraction.dispatcher import extract
    from app.models.document import DocumentStatus
    from app.services.ingest import ingest

    inner = fitz.open()
    ip = inner.new_page()
    ip.insert_text((20, 40), "Hasta: Selin Ozdemir TCKN: 32944882886", fontsize=13)
    pix = ip.get_pixmap(dpi=60)  # testi küçük tut; ölçülen şey kaplama oranı, çözünürlük değil
    inner.close()

    d = fitz.open()
    p = d.new_page()
    # 50 karakterden UZUN üstbilgi → eski koşulun ilk yarısı düşüyor
    p.insert_text((72, 60), "SIRKET ICI YAZISMA - GIZLI - DAGITIM SINIRLIDIR - 2026", fontsize=9)
    p.insert_image(fitz.Rect(30, 80, 580, 760), pixmap=pix)  # sayfanın çoğu görüntü
    data = d.tobytes()
    d.close()

    _, content = extract(data, "scan.pdf", max_bytes=50 * 1024 * 1024)
    assert IMAGE_TEXT_UNREAD in content.warnings

    doc, _ = ingest(data, "scan.pdf", Settings())
    assert doc.status == DocumentStatus.NEEDS_HUMAN_REVIEW
    assert doc.review_reason, "gözden geçirene gerekçe gösterilmeli"
