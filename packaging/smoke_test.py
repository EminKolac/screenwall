"""End-to-end smoke test of a running, packaged Screenwall (the release zip's launcher).

Run it with the backend's environment, which has httpx and python-docx:

    cd Screenwall/app/backend && UV_PROJECT_ENVIRONMENT="$SCREENWALL_HOME/venv" \
        uv run --frozen --no-dev python /path/to/packaging/smoke_test.py --base-url http://127.0.0.1:8765

or simply "$SCREENWALL_HOME/venv/bin/python packaging/smoke_test.py" (Windows:
"%SCREENWALL_HOME%\\venv\\Scripts\\python.exe packaging\\smoke_test.py").

What it checks, through the same HTTP API the UI uses (backend/app/api/*.py, security/web.py):
  1. GET /health answers {"status": "ok"} (waits up to --wait seconds).
  2. GET / is the built UI (text/html, contains "Screenwall") and its JavaScript bundle is served as
     JavaScript (a text/plain bundle is refused by the browser under nosniff).
  3. GET /api/session sets the session cookie and returns the CSRF token. With
     --launch-token-file (the launcher's SCREENWALL_HOME/launch-token), a session request without
     the token must first get 401 launch_token_required, and the session is then opened with the
     X-Launch-Token header, as the UI does with the token from its address.
  4. The local web security is on: a POST without X-CSRF-Token gets 403 csrf_failed and a foreign
     Origin gets 403 origin_not_allowed.
  5. A synthetic DOCX (obviously fake values; the T.C. kimlik number is generated with a valid
     checksum) is uploaded with the Origin and X-CSRF-Token headers, as the browser sends them.
  6. The document leaves the processing states (queued, extracting, masking, auditing) and lands
     in needs_review or approved, with counts.masked >= 1.
  7. The masked text contains at least one <TYPE_n> placeholder and none of the fake values, and
     the review payload lists the masked spans.
  Options: --expect-auditor unavailable|available asserts the auditor state (CI runs without
  Ollama: the document must wait in review with the auditor_unavailable reason). --approve tries
  to approve a document left in review, like the reviewer's button (a second click after
  audit_not_clean), and downloads the approved PDF. --require-approved fails unless it ends
  approved.

Output: one JSON summary on stdout (progress goes to stderr). Exit 0 when every check passed,
1 when one failed, 2 for bad arguments. Nothing from the document is printed except placeholder
types and counts. The uploaded test document stays in the app's store (nothing is deleted).
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import time
from typing import Any

try:
    import httpx
    from docx import Document as DocxDocument
except ImportError as missing:  # pragma: no cover - environment problem, not a test result
    print(json.dumps({"ok": False, "error": f"missing dependency: {missing.name}; run this "
                      "with the backend's environment (see the module docstring)"}))
    sys.exit(1)

PROCESSING = {"queued", "extracting", "masking", "auditing"}
SETTLED_OK = {"needs_review", "approved"}
PLACEHOLDER = re.compile(r"<([A-Z][A-Z0-9_]*)_\d+>")
SCRIPT_SRC = re.compile(r"""<script[^>]+src=["'](/assets/[^"']+\.js)["']""")
FOREIGN_ORIGIN = "http://smoke-test.invalid"

# ---------------------------------------------------------------- synthetic test data


def fake_tckn(first_nine: str = "100000000") -> str:
    """A T.C. kimlik number with valid check digits (app/anonymization/recognizers/turkish.py
    `valid_tckn`), built from an obviously fake prefix."""
    d = [int(c) for c in first_nine]
    tenth = ((d[0] + d[2] + d[4] + d[6] + d[8]) * 7 - (d[1] + d[3] + d[5] + d[7])) % 10
    eleventh = (sum(d) + tenth) % 10
    return f"{first_nine}{tenth}{eleventh}"


def fake_tr_iban(bank: str = "00099", account: str = "0000000000000001") -> str:
    """A Turkish IBAN with valid ISO 13616 check digits for an obviously fake account."""
    bban = f"{bank}0{account}"
    check = 98 - int(f"{bban}292700") % 97  # "TR" = 29 27, check digits "00" while computing
    raw = f"TR{check:02d}{bban}"
    return " ".join(raw[i:i + 4] for i in range(0, len(raw), 4))


FAKE_NAME = "DENEME ÖRNEKOĞLU"
FAKE_TCKN = fake_tckn()
FAKE_EMAIL = "deneme.ornekoglu@example.com"
FAKE_IBAN = fake_tr_iban()


def synthetic_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Screenwall duman testi", level=1)
    doc.add_paragraph("Bu belge otomatik test için üretilmiştir; tüm değerler uydurmadır.")
    doc.add_paragraph(f"Adı Soyadı: {FAKE_NAME}")
    doc.add_paragraph(f"T.C. Kimlik No: {FAKE_TCKN}")
    doc.add_paragraph(f"E-posta: {FAKE_EMAIL}")
    doc.add_paragraph(f"IBAN: {FAKE_IBAN}")
    doc.add_paragraph("Bu bilgiler yalnızca maskeleme denemesi içindir.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def fake_values() -> list[str]:
    """The raw values that must not survive masking (the IBAN with and without spaces)."""
    return [FAKE_TCKN, FAKE_EMAIL, FAKE_IBAN, FAKE_IBAN.replace(" ", "")]


# ---------------------------------------------------------------- reporting


class Report:
    def __init__(self, base_url: str) -> None:
        self.started = time.monotonic()
        self.data: dict[str, Any] = {"ok": True, "base_url": base_url, "checks": []}

    def check(self, name: str, ok: bool, detail: Any = None) -> bool:
        entry: dict[str, Any] = {"name": name, "ok": bool(ok)}
        if detail is not None:
            entry["detail"] = detail
        self.data["checks"].append(entry)
        if not ok:
            self.data["ok"] = False
        log(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f": {detail}" if detail is not None else ""))
        return bool(ok)

    def finish(self) -> int:
        self.data["elapsed_s"] = round(time.monotonic() - self.started, 1)
        # ASCII escapes: a Windows CI console may not encode UTF-8 (cp1252 when redirected).
        print(json.dumps(self.data, indent=2))
        return 0 if self.data["ok"] else 1


def log(message: str) -> None:
    print(message, file=sys.stderr, flush=True)


def detail_of(response: httpx.Response) -> str:
    """The API's {"detail": <ErrorCode>} or the status line."""
    try:
        body = response.json()
        if isinstance(body, dict) and isinstance(body.get("detail"), str):
            return body["detail"]
    except ValueError:
        pass
    return f"HTTP {response.status_code}"


# ---------------------------------------------------------------- steps


def wait_for_health(client: httpx.Client, report: Report, wait_s: float) -> bool:
    deadline = time.monotonic() + wait_s
    last = "no answer"
    while True:
        try:
            response = client.get("/health", timeout=5)
            if response.status_code == 200 and response.json().get("status") == "ok":
                return report.check("health", True)
            last = f"HTTP {response.status_code}"
        except (httpx.HTTPError, ValueError) as error:
            last = type(error).__name__
        if time.monotonic() >= deadline:
            return report.check("health", False, f"/health not ok after {wait_s:.0f} s ({last})")
        time.sleep(1)


def check_ui(client: httpx.Client, report: Report) -> None:
    response = client.get("/")
    kind = response.headers.get("content-type", "")
    ok = response.status_code == 200 and kind.startswith("text/html") \
        and "Screenwall" in response.text
    report.check("ui_index", ok, f"HTTP {response.status_code}, {kind}")
    scripts = SCRIPT_SRC.findall(response.text) if ok else []
    if not scripts:
        report.check("ui_bundle", False, "no /assets/*.js script in index.html")
        return
    bundle = client.get(scripts[0])
    kind = bundle.headers.get("content-type", "")
    report.check("ui_bundle", bundle.status_code == 200 and "javascript" in kind,
                 f"HTTP {bundle.status_code}, {kind}")


def read_launch_token(path: str, report: Report) -> str | None:
    try:
        with open(path, encoding="ascii") as handle:
            token = handle.readline().strip()
    except (OSError, UnicodeError) as error:
        report.check("launch_token_file", False, f"not readable: {type(error).__name__}")
        return None
    report.check("launch_token_file", bool(token), None if token else "empty")
    return token or None


def check_launch_token_required(client: httpx.Client, report: Report) -> None:
    """Another account on the computer reaches 127.0.0.1 but cannot read the token file."""
    response = client.get("/api/session")
    report.check("launch_token_required",
                 response.status_code == 401 and detail_of(response) == "launch_token_required"
                 and not client.cookies, f"{response.status_code} {detail_of(response)}")
    client.cookies.clear()


def open_session(client: httpx.Client, report: Report, launch_token: str | None) -> str | None:
    headers = {"X-Launch-Token": launch_token} if launch_token else {}
    response = client.get("/api/session", headers=headers)
    token = response.json().get("csrf_token") if response.status_code == 200 else None
    report.check("session", bool(token) and bool(client.cookies),
                 None if token else detail_of(response))
    return token


def check_api_health(client: httpx.Client, report: Report, expect: str) -> dict:
    response = client.get("/api/health")
    if not report.check("api_health", response.status_code == 200, detail_of(response)
                        if response.status_code != 200 else None):
        return {}
    health = response.json()
    report.data["health"] = health
    report.check("worker_alive", bool(health.get("worker", {}).get("alive")))
    reachable = bool(health.get("auditor", {}).get("reachable"))
    if expect == "available":
        report.check("auditor_reachable", reachable, health.get("auditor"))
    elif expect == "unavailable":
        report.check("auditor_unreachable", not reachable, health.get("auditor"))
    return health


def check_security(client: httpx.Client, report: Report, csrf: str) -> None:
    files = {"file": ("x.docx", b"not a document", "application/octet-stream")}
    no_csrf = client.post("/api/documents", files=files)
    report.check("csrf_required",
                 no_csrf.status_code == 403 and detail_of(no_csrf) == "csrf_failed",
                 f"{no_csrf.status_code} {detail_of(no_csrf)}")
    foreign = client.post("/api/documents", files=files,
                          headers={"Origin": FOREIGN_ORIGIN, "X-CSRF-Token": csrf})
    report.check("foreign_origin_refused",
                 foreign.status_code == 403 and detail_of(foreign) == "origin_not_allowed",
                 f"{foreign.status_code} {detail_of(foreign)}")


def upload(client: httpx.Client, report: Report, csrf: str) -> dict | None:
    files = {"file": ("screenwall-smoke-test.docx", synthetic_docx(),
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/documents", files=files, headers={"X-CSRF-Token": csrf})
    ok = response.status_code == 202
    summary = response.json() if ok else None
    report.check("upload", ok and bool(summary and summary.get("id")),
                 f"HTTP {response.status_code}" + ("" if ok else f" {detail_of(response)}"))
    return summary


def wait_until_settled(client: httpx.Client, report: Report, doc_id: str,
                       timeout_s: float) -> dict | None:
    started = time.monotonic()
    seen: list[str] = []
    summary: dict | None = None
    while time.monotonic() - started < timeout_s:
        response = client.get(f"/api/documents/{doc_id}")
        if response.status_code != 200:
            report.check("processing", False, f"HTTP {response.status_code} {detail_of(response)}")
            return None
        summary = response.json()
        status = summary.get("status")
        if not seen or seen[-1] != status:
            seen.append(status)
            log(f"  status: {status}")
        if status not in PROCESSING:
            seconds = round(time.monotonic() - started, 1)
            report.check("processing", True, {"seconds": seconds, "states": seen})
            return summary
        time.sleep(1)
    report.check("processing", False, f"still {seen[-1] if seen else '?'} after {timeout_s:.0f} s")
    return None


def describe(summary: dict) -> dict:
    return {
        "id": summary.get("id"),
        "status": summary.get("status"),
        "mode": summary.get("mode"),
        "language": summary.get("language"),
        "revision": summary.get("revision"),
        "masked": summary.get("counts", {}).get("masked"),
        "pending": summary.get("counts", {}).get("pending"),
        "reasons": [r.get("code") for r in summary.get("reasons", [])],
        "can": summary.get("can"),
    }


def check_masking(client: httpx.Client, report: Report, summary: dict) -> None:
    doc_id = summary["id"]
    masked_count = summary.get("counts", {}).get("masked") or 0
    report.check("masked_count", masked_count >= 1, masked_count)

    response = client.get(f"/api/documents/{doc_id}/masked")
    if not report.check("masked_text", response.status_code == 200,
                        None if response.status_code == 200 else detail_of(response)):
        return
    text = response.text
    types: dict[str, int] = {}
    for kind in PLACEHOLDER.findall(text):
        types[kind] = types.get(kind, 0) + 1
    report.data["placeholders"] = types
    report.check("placeholder_present", bool(types), types or "no <TYPE_n> placeholder")
    leaked = [i for i, value in enumerate(fake_values()) if value in text]
    report.check("fake_values_masked", not leaked,
                 None if not leaked else f"{len(leaked)} fake value(s) left in the masked text")

    review = client.get(f"/api/documents/{doc_id}/review")
    if report.check("review_payload", review.status_code == 200,
                    None if review.status_code == 200 else detail_of(review)):
        spans = review.json().get("spans", [])
        masked_spans = [s for s in spans if s.get("state") in ("masked", "added")]
        report.check("review_spans", bool(masked_spans),
                     {"masked_spans": len(masked_spans), "placeholders":
                      sorted({s.get("placeholder") for s in masked_spans} - {""})})


def try_approve(client: httpx.Client, report: Report, csrf: str, summary: dict) -> dict:
    """Approve like the reviewer's button: POST approve with the revision seen; after
    409 audit_not_clean the reasons are on the document and a second click approves."""
    attempts: list[str] = []
    for _ in range(2):
        if summary.get("status") != "needs_review":
            break
        response = client.post(f"/api/documents/{summary['id']}/approve",
                               json={"expected_revision": summary.get("revision")},
                               headers={"X-CSRF-Token": csrf}, timeout=600)
        attempts.append(f"{response.status_code} {detail_of(response)}"
                        if response.status_code != 200 else "200")
        if response.status_code == 200:
            summary = response.json()
            break
        if detail_of(response) != "audit_not_clean":
            break
        summary = client.get(f"/api/documents/{summary['id']}").json()
    report.data["approval"] = {"attempts": attempts}
    return summary


def check_download(client: httpx.Client, report: Report, doc_id: str) -> None:
    response = client.get(f"/api/documents/{doc_id}/download", timeout=120)
    kind = response.headers.get("content-type", "")
    ok = response.status_code == 200 and kind.startswith("application/pdf") \
        and response.content.startswith(b"%PDF")
    report.check("download_pdf", ok, f"HTTP {response.status_code}, {kind}, "
                 f"{len(response.content)} bytes")
    if not ok:
        return
    try:
        import fitz  # PyMuPDF, a backend dependency

        with fitz.open(stream=response.content, filetype="pdf") as pdf:
            text = "".join(page.get_text() for page in pdf)
    except Exception as error:  # noqa: BLE001 - the PDF check is best effort
        report.check("download_text", False, f"could not read the PDF: {type(error).__name__}")
        return
    leaked = [v for v in fake_values() if v in text]
    report.check("download_text", bool(PLACEHOLDER.search(text)) and not leaked,
                 {"placeholder": bool(PLACEHOLDER.search(text)), "fake_values_left": len(leaked)})


# ---------------------------------------------------------------- main


def run(args: argparse.Namespace) -> int:
    base = args.base_url.rstrip("/")
    report = Report(base)
    # Every request carries the UI's own Origin, as the browser's fetch calls do.
    with httpx.Client(base_url=base, timeout=60, headers={"Origin": base},
                      follow_redirects=False) as client:
        try:
            if not wait_for_health(client, report, args.wait):
                return report.finish()
            check_ui(client, report)
            launch_token = None
            if args.launch_token_file:
                launch_token = read_launch_token(args.launch_token_file, report)
                if not launch_token:
                    return report.finish()
                check_launch_token_required(client, report)
            csrf = open_session(client, report, launch_token)
            if not csrf:
                return report.finish()
            check_api_health(client, report, args.expect_auditor)
            check_security(client, report, csrf)
            uploaded = upload(client, report, csrf)
            if not uploaded:
                return report.finish()
            summary = wait_until_settled(client, report, uploaded["id"], args.timeout)
            if summary is None:
                return report.finish()
            report.data["document"] = describe(summary)
            report.check("settled", summary.get("status") in SETTLED_OK, summary.get("status"))
            if summary.get("status") not in SETTLED_OK:
                return report.finish()
            if args.expect_auditor == "unavailable":
                report.check("auditor_unavailable_reason",
                             summary.get("status") == "needs_review" and "auditor_unavailable"
                             in report.data["document"]["reasons"],
                             report.data["document"]["reasons"])
            check_masking(client, report, summary)
            if args.approve and summary.get("status") == "needs_review":
                summary = try_approve(client, report, csrf, summary)
                report.data["document_after_approval"] = describe(summary)
            approved = summary.get("status") == "approved"
            report.data["approved"] = approved
            if approved:
                check_download(client, report, summary["id"])
            if args.require_approved:
                report.check("approved", approved, summary.get("status"))
        except httpx.HTTPError as error:
            report.check("http", False, f"{type(error).__name__}: {error}")
    return report.finish()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--base-url", default="http://127.0.0.1:8765")
    parser.add_argument("--launch-token-file", default=None,
                        help="the launcher's SCREENWALL_HOME/launch-token: check that a session "
                             "needs it, then open one with it")
    parser.add_argument("--wait", type=float, default=60,
                        help="seconds to wait for /health (default 60)")
    parser.add_argument("--timeout", type=float, default=300,
                        help="seconds to wait for processing to end (default 300)")
    parser.add_argument("--expect-auditor", choices=("any", "available", "unavailable"),
                        default="any", help="assert the auditor state (default any)")
    parser.add_argument("--approve", action="store_true",
                        help="approve a document left in review, then download it")
    parser.add_argument("--require-approved", action="store_true",
                        help="fail unless the document ends approved")
    return run(parser.parse_args())


if __name__ == "__main__":
    sys.exit(main())
